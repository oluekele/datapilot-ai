import io
import json
import re
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from datapilot import worker

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "customer_churn.csv"
CONTRACT_KINDS = {
    "numeric", "categorical", "boolean", "datetime", "text", "identifier", "empty",
}
SAFE_NAME = re.compile(r"^[a-z0-9][a-z0-9_.-]*$")
PNG_MAGIC = b"\x89PNG\r\n\x1a\n"


def run(monkeypatch, job):
    monkeypatch.setattr("sys.stdin", io.StringIO(json.dumps(job)))
    out = io.StringIO()
    monkeypatch.setattr("sys.stdout", out)
    assert worker.main() == 0
    return json.loads(out.getvalue())


def job(op, params, tmp_path, input_path=EXAMPLE):
    return {
        "type": op, "input_path": str(input_path), "filename": "c.csv",
        "dataset_name": "Churn", "output_dir": str(tmp_path / "out"),
        "params": params,
    }


def check_artifacts(out, tmp_path):
    """Every artifact exists in output_dir under a safe, unique name."""
    directory = tmp_path / "out"
    names = [a["file"] for a in out["artifacts"]]
    assert len(names) == len(set(names))
    assert sorted(names) == sorted(p.name for p in directory.iterdir())
    for artifact in out["artifacts"]:
        assert SAFE_NAME.match(artifact["file"]), artifact["file"]
        assert artifact["kind"] in {"CHART", "TABLE", "MODEL", "REPORT", "CODE", "DATA"}
        assert artifact["title"]
        assert isinstance(artifact["metadata"], dict)
        path = directory / artifact["file"]
        assert path.stat().st_size > 0
        if artifact["content_type"] == "image/png":
            assert path.read_bytes()[:8] == PNG_MAGIC
    return {a["file"]: a for a in out["artifacts"]}


# ---------------------------------------------------------------- PROFILE


def test_profile_reports_column_kinds_and_no_artifacts(monkeypatch):
    out = run(monkeypatch, {"type": "PROFILE", "input_path": str(EXAMPLE),
                            "filename": "c.csv"})
    assert out["ok"] and out["artifacts"] == []
    kinds = out["result"]["columnKinds"]
    assert set(kinds.values()) <= CONTRACT_KINDS
    assert kinds["customer_id"] == "identifier"
    assert kinds["churned"] == "boolean"
    assert kinds["monthly_charge"] == "numeric"
    assert kinds["plan"] == "categorical"


def test_column_kinds_detects_empty_and_text(tmp_path):
    from datapilot.operations import column_kinds

    df = pd.DataFrame({
        "blank": [None] * 60,
        "comment": [f"note {i % 55}" for i in range(60)],
        "order_id": range(60),
        "when": pd.date_range("2024-01-01", periods=60),
    })
    kinds = column_kinds(df)
    assert kinds == {"blank": "empty", "comment": "text", "order_id": "identifier",
                     "when": "datetime"}


# ---------------------------------------------------------------- EXPLORE


def test_explore_default(monkeypatch, tmp_path):
    out = run(monkeypatch, job("EXPLORE", {}, tmp_path))
    assert out["ok"], out
    result = out["result"]
    assert set(result) == {"summary", "missing", "distributions", "categories",
                           "correlations"}
    assert result["summary"]["rows"] == 403
    assert {d["column"] for d in result["distributions"]} == {
        "tenure_months", "monthly_charge", "support_tickets", "satisfaction_score"}
    assert all("mean" in d["stats"] for d in result["distributions"])
    plan = next(c for c in result["categories"] if c["column"] == "plan")
    assert {"value", "count", "percentage"} <= set(plan["frequencies"][0])
    assert result["correlations"]["method"] == "pearson"
    assert {"column_a", "column_b", "correlation", "strength"} <= set(
        result["correlations"]["top"][0])
    assert result["missing"]["columns_with_missing"] == 3

    files = check_artifacts(out, tmp_path)
    assert files["histogram-monthly_charge.png"]["kind"] == "CHART"
    assert "bar-plan.png" in files
    assert "correlation-heatmap-pearson.png" in files
    assert "missing-values.png" in files
    assert files["eda-report.md"]["content_type"] == "text/markdown"
    workbook = files["eda-workbook.xlsx"]
    assert workbook["content_type"].endswith("spreadsheetml.sheet")
    sheets = load_workbook(tmp_path / "out" / "eda-workbook.xlsx").sheetnames
    assert {"Summary", "Missing", "Distributions", "Frequencies",
            "Correlations"} <= set(sheets)
    report = (tmp_path / "out" / "eda-report.md").read_text()
    assert "# Exploratory Data Analysis" in report and "Churn" in report


def test_explore_selected_columns_spearman(monkeypatch, tmp_path):
    params = {"correlationMethod": "spearman", "columns": ["monthly_charge", "plan"]}
    out = run(monkeypatch, job("EXPLORE", params, tmp_path))
    assert out["ok"], out
    assert [d["column"] for d in out["result"]["distributions"]] == ["monthly_charge"]
    assert [c["column"] for c in out["result"]["categories"]] == ["plan"]
    assert out["result"]["correlations"]["method"] == "spearman"
    files = check_artifacts(out, tmp_path)
    assert "histogram-tenure_months.png" not in files


# ---------------------------------------------------------------- ANALYZE


@pytest.mark.parametrize(
    ("params", "chart"),
    [
        ({"test": "descriptive"}, False),
        ({"test": "descriptive", "columns": ["tenure_months"]}, False),
        ({"test": "grouped_summary", "groupColumn": "plan",
          "valueColumn": "monthly_charge"}, True),
        ({"test": "compare_groups", "groupColumn": "churned",
          "valueColumn": "tenure_months", "alpha": 0.01}, True),
        ({"test": "association", "columnA": "plan", "columnB": "city"}, True),
        ({"test": "correlation", "columnX": "tenure_months",
          "columnY": "monthly_charge", "method": "spearman"}, True),
    ],
)
def test_analyze(monkeypatch, tmp_path, params, chart):
    out = run(monkeypatch, job("ANALYZE", params, tmp_path))
    assert out["ok"], out
    result = out["result"]
    assert result["test"] == params["test"]
    assert result["title"] and isinstance(result["details"], dict)
    assert isinstance(result["summary"], str) and len(result["summary"]) > 30
    assert isinstance(result["table"], list) and result["table"]
    files = check_artifacts(out, tmp_path)
    types = [a["content_type"] for a in files.values()]
    assert types.count("text/csv") == 1
    assert types.count("text/markdown") == 1
    assert ("image/png" in types) is chart


def test_compare_groups_details(monkeypatch, tmp_path):
    params = {"test": "compare_groups", "groupColumn": "churned",
              "valueColumn": "tenure_months"}
    result = run(monkeypatch, job("ANALYZE", params, tmp_path))["result"]
    assert result["details"]["primary_test"] in ("Welch's t-test",
                                                 "Mann–Whitney U test")
    assert result["details"]["significant"] is True
    assert "causes" in result["summary"]


# ---------------------------------------------------------------- VISUALIZE


def test_visualize_png_and_svg(monkeypatch, tmp_path):
    params = {
        "charts": [
            {"chartType": "scatter", "x": "tenure_months", "y": "monthly_charge",
             "group": "plan", "filters": [{"column": "city",
                                           "values": ["Abuja", "Lagos"]}]},
            {"chartType": "bar", "x": "plan", "y": "monthly_charge",
             "aggregation": "mean", "title": "Average charge per plan"},
            {"chartType": "histogram", "x": "tenure_months", "bins": 15,
             "filters": [{"column": "monthly_charge", "minValue": 10,
                          "maxValue": 60}]},
        ],
        "formats": ["png", "svg"],
    }
    out = run(monkeypatch, job("VISUALIZE", params, tmp_path))
    assert out["ok"], out
    charts = out["result"]["charts"]
    assert [c["title"] for c in charts][1] == "Average charge per plan"
    assert "filter city in [Abuja, Lagos]" in charts[0]["description"]
    files = check_artifacts(out, tmp_path)
    assert len(files) == 6
    svgs = [n for n, a in files.items() if a["content_type"] == "image/svg+xml"]
    assert len(svgs) == 3
    assert b"<svg" in (tmp_path / "out" / svgs[0]).read_bytes()[:500]
    assert all(a["kind"] == "CHART" for a in files.values())


def test_visualize_defaults_to_png(monkeypatch, tmp_path):
    params = {"charts": [{"chartType": "correlation_heatmap"}]}
    out = run(monkeypatch, job("VISUALIZE", params, tmp_path))
    assert out["ok"], out
    assert [a["content_type"] for a in out["artifacts"]] == ["image/png"]


# ---------------------------------------------------------------- user errors


@pytest.mark.parametrize(
    ("op", "params", "message"),
    [
        ("EXPLORE", {"correlationMethod": "kendall"}, "correlationMethod"),
        ("EXPLORE", {"columns": ["customer_id"]}, "identifier"),
        ("EXPLORE", {"columns": ["nope"]}, "does not exist"),
        ("EXPLORE", {"columns": "plan"}, "must be a list"),
        ("EXPLORE", ["not", "an", "object"], "JSON object"),
        ("ANALYZE", {}, "'test' is required"),
        ("ANALYZE", {"test": "anova"}, "Unsupported test"),
        ("ANALYZE", {"test": "grouped_summary", "groupColumn": "plan"},
         "'valueColumn' is required"),
        ("ANALYZE", {"test": "grouped_summary", "groupColumn": "customer_id",
                     "valueColumn": "monthly_charge"}, "at most 50 groups"),
        ("ANALYZE", {"test": "compare_groups", "groupColumn": "plan",
                     "valueColumn": "city"}, "must be numeric"),
        ("ANALYZE", {"test": "association", "columnA": "plan",
                     "columnB": "monthly_charge"}, "categorical"),
        ("ANALYZE", {"test": "correlation", "columnX": "tenure_months",
                     "columnY": "monthly_charge", "alpha": 2}, "alpha"),
        ("ANALYZE", {"test": "correlation", "columnX": "tenure_months",
                     "columnY": "monthly_charge", "alpha": True}, "number"),
        ("VISUALIZE", {}, "'charts' must be a list"),
        ("VISUALIZE", {"charts": []}, "between 1 and 12"),
        ("VISUALIZE", {"charts": [{"chartType": "bar", "x": "plan"}] * 13},
         "between 1 and 12"),
        ("VISUALIZE", {"charts": [{"chartType": "pie", "x": "plan"}]},
         "Chart 1: Unsupported chartType"),
        ("VISUALIZE", {"charts": [{"chartType": "histogram", "x": "plan"}]},
         "numeric x column"),
        ("VISUALIZE", {"charts": [{"chartType": "bar", "x": "plan",
                                   "filters": [{"column": "ghost", "values": [1]}]}]},
         "does not exist"),
        ("VISUALIZE", {"charts": [{"chartType": "bar", "x": "plan",
                                   "filters": [{"column": "city",
                                                "values": ["Atlantis"]}]}]},
         "no rows remain"),
        ("VISUALIZE", {"charts": [{"chartType": "bar", "x": "plan"}],
                       "formats": ["pdf"]}, "Unsupported formats"),
        ("VISUALIZE", {"charts": [{"chartType": "histogram", "x": "tenure_months",
                                   "bins": 1}]}, "bins"),
    ],
)
def test_user_errors(monkeypatch, tmp_path, op, params, message):
    out = run(monkeypatch, job(op, params, tmp_path))
    assert out["ok"] is False
    assert message in out["error"]


def test_missing_output_dir_is_reported(monkeypatch, tmp_path):
    payload = job("EXPLORE", {}, tmp_path)
    del payload["output_dir"]
    out = run(monkeypatch, payload)
    assert out == {"ok": False, "error": "The job has no output directory."}
